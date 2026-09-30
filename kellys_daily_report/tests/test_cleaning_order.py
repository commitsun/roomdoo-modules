import base64
from io import BytesIO

from openpyxl import load_workbook

from odoo.addons.pms.tests.common import TestPms


class TestCleaningOrder(TestPms):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.room_type = cls.env["pms.room.type"].create(
            {
                "pms_property_ids": [cls.pms_property1.id],
                "name": "Double Test",
                "default_code": "DBL_Test",
                "class_id": cls.room_type_class1.id,
            }
        )
        cls.rooms = cls.env["pms.room"]
        for sequence, name in enumerate(["10", "2", "B", "A"]):
            cls.rooms |= cls.env["pms.room"].create(
                {
                    "pms_property_id": cls.pms_property1.id,
                    "name": name,
                    "room_type_id": cls.room_type.id,
                    "capacity": 2,
                    "sequence": sequence,
                }
            )

    def _create_wizard(self):
        kellys_rooms = self.env["kellysrooms"]
        for room in self.rooms:
            kellys_rooms |= self.env["kellysrooms"].create(
                {
                    "habitacion": room.name,
                    "habitacionid": room.id,
                    "tipo": "1",
                }
            )
        return self.env["kellysreport"].create(
            {
                "pms_property_id": self.pms_property1.id,
                "habitaciones": [(6, 0, kellys_rooms.ids)],
            }
        )

    def _set_cleaning_order(self, order_by_name):
        for room in self.rooms:
            room.cleaning_sequence = order_by_name[room.name]

    def test_cleaning_sequence_defaults_to_zero(self):
        """New rooms get no cleaning order, whatever their sequence."""
        self.assertEqual(self.rooms.mapped("cleaning_sequence"), [0, 0, 0, 0])

    def test_rooms_without_cleaning_order_are_sorted_by_name(self):
        """Without cleaning order the listing keeps the order by room name."""
        wizard = self._create_wizard()
        self.assertEqual(
            wizard._get_sorted_rooms().mapped("habitacion"), ["10", "2", "A", "B"]
        )

    def test_rooms_are_sorted_by_cleaning_order(self):
        """The cleaning order wins over the room name and the room sequence."""
        self._set_cleaning_order({"10": 3, "2": 1, "B": 2, "A": 4})
        wizard = self._create_wizard()
        self.assertEqual(
            wizard._get_sorted_rooms().mapped("habitacion"), ["2", "B", "10", "A"]
        )

    def test_rooms_with_same_cleaning_order_are_sorted_by_name(self):
        """Rooms that share a cleaning order are listed by name."""
        self._set_cleaning_order({"10": 1, "2": 1, "B": 0, "A": 0})
        wizard = self._create_wizard()
        self.assertEqual(
            wizard._get_sorted_rooms().mapped("habitacion"), ["A", "B", "10", "2"]
        )

    def test_excel_export_follows_cleaning_order(self):
        """The Excel export (used by the front) lists rooms in cleaning order."""
        self._set_cleaning_order({"10": 3, "2": 1, "B": 2, "A": 4})
        wizard = self._create_wizard()
        result = wizard._excel_export()
        workbook = load_workbook(BytesIO(base64.decodebytes(result["xls_binary"])))
        worksheet = workbook.active
        room_names = [
            str(row[0]) for row in worksheet.iter_rows(min_row=2, values_only=True)
        ]
        self.assertEqual(room_names, ["2", "B", "10", "A"])

    def test_pdf_report_follows_cleaning_order(self):
        """The PDF report receives the rooms in cleaning order."""
        self._set_cleaning_order({"10": 3, "2": 1, "B": 2, "A": 4})
        wizard = self._create_wizard()
        action = wizard.with_context(discard_logo_check=True).print_rooms_report()
        rooms = self.env["kellysrooms"].browse(action["context"]["active_ids"])
        self.assertEqual(rooms.mapped("habitacion"), ["2", "B", "10", "A"])
